#!/usr/bin/env python3
"""Generate AgriCheck V3 dev reference files: users-per-agency Excel + workflow PDF."""

from __future__ import annotations

import json
import os
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "dev-reference"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Known dev passwords from DatabaseSeeder / docs/PROGRESS.md
SEED_USERS = [
    # System
    ("SYSTEM", "System Admin", "admin@agricheck.local", "Admin@12345", "ROLE_ADMIN", "System Administrator", "/admin", "Global system admin — payment config for all agencies"),
    # Client
    ("CLIENT", "Importer (Client)", "importer@agricheck.local", "Importer@12345", "ROLE_IMPORTER", "Demo Importer", "/client", "Demo client / importer"),
    # BAI
    ("BAI", "Bureau of Animal Industry", "evaluator@agricheck.local", "Evaluator@12345", "ROLE_EVALUATOR", "BAI Evaluator", "/agency", "Seeded"),
    ("BAI", "Bureau of Animal Industry", "inspector@agricheck.local", "Inspector@12345", "ROLE_INSPECTOR", "BAI Inspector", "/inspector", "Seeded — shared inspector account"),
    ("BAI", "Bureau of Animal Industry", "billing@agricheck.local", "Billing@12345", "ROLE_BILLING_AGENT", "BAI Billing Agent", "/agency", "Seeded — DA billing verification"),
    # BFAR
    ("BFAR", "Bureau of Fisheries and Aquatic Resources", "evaluator.bfar@agricheck.local", "Evaluator@12345", "ROLE_EVALUATOR", "BFAR Evaluator", "/agency", "Seeded"),
    ("BPI", "Bureau of Plant Industry", "evaluator.bpi@agricheck.local", "Evaluator@12345", "ROLE_EVALUATOR", "BPI Evaluator", "/agency", "Seeded"),
    ("BAI", "Bureau of Animal Industry", "bai.admin@agricheck.local", "AgencyAdmin@12345", "ROLE_AGENCY_ADMIN", "BAI Agency Admin", "/agency", "Seeded - agency payment config"),
    ("BFAR", "Bureau of Fisheries and Aquatic Resources", "bfar.admin@agricheck.local", "AgencyAdmin@12345", "ROLE_AGENCY_ADMIN", "BFAR Agency Admin", "/agency", "Seeded - agency payment config"),
    ("BPI", "Bureau of Plant Industry", "bpi.admin@agricheck.local", "AgencyAdmin@12345", "ROLE_AGENCY_ADMIN", "BPI Agency Admin", "/agency", "Seeded - agency payment config"),
    ("SRA", "Sugar Regulatory Administration", "sra.admin@agricheck.local", "AgencyAdmin@12345", "ROLE_AGENCY_ADMIN", "SRA Agency Admin", "/agency", "Seeded - agency payment config"),
    ("NTA", "National Tobacco Administration", "nta.admin@agricheck.local", "AgencyAdmin@12345", "ROLE_AGENCY_ADMIN", "NTA Agency Admin", "/agency", "Seeded - agency payment config"),
    # DA
    ("DA", "Department of Agriculture", "accred@agricheck.local", "Accred@12345", "ROLE_ACCREDITATION_OFFICER", "DA Accreditation Officer", "/agency", "Seeded"),
    ("DA", "Department of Agriculture", "daevaluator@agricheck.local", "DaEval@12345", "ROLE_ACCREDITATION_OFFICER", "DA Accreditation Evaluator", "/agency", "Seeded"),
    ("DA", "Department of Agriculture", "da.secretary@agricheck.local", "DaSec@12345", "ROLE_DA_SECRETARY", "DA Secretary", "/da", "Seeded"),
    ("DA", "Department of Agriculture", "da.undersecretary@agricheck.local", "DaUsec@12345", "ROLE_DA_UNDERSECRETARY", "DA Undersecretary", "/da", "Seeded"),
    # MAV
    ("MAV", "MAV Management Committee", "mav.admin@agricheck.local", "MavAdmin@12345", "ROLE_MAV_ADMIN", "MAV Administrator", "/mav/admin", "Seeded"),
    ("MAV", "MAV Management Committee", "mav.evaluator@agricheck.local", "MavEval@12345", "ROLE_MAV_EVALUATOR", "MAV Evaluator", "/mav/admin", "Seeded"),
    ("MAV", "MAV Management Committee", "mav.secretary@agricheck.local", "MavSec@12345", "ROLE_MAV_SECRETARY", "MAV Secretary", "/mav/admin", "Seeded"),
    # Ops (cross-agency)
    ("OPS", "Operations (Cross-Agency)", "warehouse@agricheck.local", "Warehouse@12345", "ROLE_WAREHOUSE_STAFF", "Warehouse Staff", "/warehouse", "Seeded"),
    ("OPS", "Operations (Cross-Agency)", "driver@agricheck.local", "Driver@12345", "ROLE_DRIVER", "Demo Driver", "/driver", "Seeded — AgriTrack mobile"),
    ("OPS", "Operations (Cross-Agency)", "operator@agricheck.local", "Operator@12345", "ROLE_OPERATOR", "AgriTrack Operator", "/driver", "Seeded — AgriTrack mobile"),
    ("OPS", "Operations (Cross-Agency)", "doctor@agricheck.local", "Doctor@12345", "ROLE_DOCTOR", "Port Doctor", "/warehouse", "Seeded"),
]

PAYMENT_CONFIG = [
    ("BAI", "Import", 2500, "PHP", True, "Seeded in processing_fee_configs"),
    ("BAI", "Export", 1500, "PHP", True, "Seeded in processing_fee_configs"),
    ("BFAR", "Import", 2000, "PHP", True, "Seeded in processing_fee_configs"),
    ("BPI", "Import", 2500, "PHP", True, "Seeded in processing_fee_configs"),
    ("SRA", "Import", 2500, "PHP", True, "Seeded in processing_fee_configs"),
    ("NTA", "Import", 2500, "PHP", True, "Seeded in processing_fee_configs"),
    ("GLOBAL", "Import", 2500, "PHP", True, "system_settings: entry_processing_fee_import"),
    ("GLOBAL", "Export", 2500, "PHP", True, "system_settings: entry_processing_fee_export"),
]

WORKFLOW_STEPS = [
    ("1", "Client", "Submit Entry + Processing Fee", "Submitted", "/client/entries/{uuid}"),
    ("2", "Evaluator", "Review Documents + Compliance Checklist", "UnderReview", "/agency/evaluator/entries/{uuid}"),
    ("3", "Evaluator", "Submit Entry Outcome - Approved", "DaIssueBilling", "/agency/evaluator/entries/{uuid}"),
    ("4", "System", "Auto-create & Issue DA Regulatory Bill", "DaIssueBilling", "billings table"),
    ("5", "Client", "Pay DA Bill + Upload Payment Proof", "DaIssueBilling", "/client/entries/{uuid} - DA Billing tab"),
    ("6", "Billing Agent", "Verify DA Payment", "ForInspection", "/agency/billing"),
    ("7", "System", "Auto-issue Import Certificate", "ForInspection", "Certificate issued on DA billing paid"),
    ("8", "Inspector", "Upload Container Photos at BOC (6 types)", "ForInspection", "/inspector/inspections"),
    ("9", "Inspector", "Approve All Container Photos", "ReadyForTransport", "/inspector/inspections"),
    ("10", "Agency/Ops", "Tag Containers for AgriTrack Pickup", "AwaitingTransport", "Transport tagging"),
    ("11", "Driver/Operator", "Claim & Deliver Containers", "InTransit", "/driver (AgriTrack mobile)"),
    ("12", "Warehouse", "Receive + Doctor Inspection", "AtWarehouse", "/warehouse"),
]


def try_load_db_users() -> list[tuple] | None:
    try:
        import pymysql
    except ImportError:
        return None

    config = {
        "host": os.environ.get("MYSQL_HOST", "localhost"),
        "port": int(os.environ.get("MYSQL_PORT", "3306")),
        "user": os.environ.get("MYSQL_USER", "agricheck_v3"),
        "password": os.environ.get("MYSQL_PASSWORD", "agricheck_v3_dev"),
        "database": os.environ.get("MYSQL_DATABASE", "agricheck_v3"),
        "charset": "utf8mb4",
        "cursorclass": pymysql.cursors.DictCursor,
    }

    sql = """
        SELECT
            COALESCE(a.Code, 'UNASSIGNED') AS agency_code,
            COALESCE(a.Name, 'No Agency Membership') AS agency_name,
            u.Email AS email,
            CONCAT(COALESCE(p.FirstName, ''), ' ', COALESCE(p.LastName, '')) AS full_name,
            GROUP_CONCAT(DISTINCT r.Code ORDER BY r.Code SEPARATOR ', ') AS roles,
            u.Status AS status
        FROM users u
        LEFT JOIN user_profiles p ON p.UserId = u.Id
        LEFT JOIN user_roles ur ON ur.UserId = u.Id
        LEFT JOIN roles r ON r.Id = ur.RoleId
        LEFT JOIN agency_memberships am ON am.UserId = u.Id
        LEFT JOIN agencies a ON a.Id = am.AgencyId
        GROUP BY u.Id, a.Code, a.Name, u.Email, p.FirstName, p.LastName, u.Status
        ORDER BY COALESCE(a.Code, 'ZZZ'), u.Email
    """

    try:
        conn = pymysql.connect(**config)
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()
        conn.close()
    except Exception as exc:
        print(f"DB query skipped: {exc}")
        return None

    result = []
    for row in rows:
        result.append((
            row["agency_code"],
            row["agency_name"],
            row["email"],
            "(see seed docs)",
            row["roles"] or "",
            row["full_name"].strip(),
            "",
            f"Status: {row['status']} - from live database",
        ))
    return result


PASSWORD_LOOKUP = {row[2].lower(): row[3] for row in SEED_USERS}
PORTAL_LOOKUP = {row[2].lower(): row[6] for row in SEED_USERS}


def enrich_db_row(row: tuple) -> tuple:
    email = row[2].lower()
    password = PASSWORD_LOOKUP.get(email, "(hashed - reset via admin or see seed docs)")
    portal = row[6] or PORTAL_LOOKUP.get(email, "")
    return (*row[:3], password, *row[4:6], portal, row[7])


def merge_users() -> list[tuple]:
    db_users = try_load_db_users()
    if db_users:
        print(f"Loaded {len(db_users)} user rows from database")
        return [enrich_db_row(row) for row in db_users]

    print("Using seeded dev user list (database unavailable)")
    return list(SEED_USERS)


def style_header(ws, headers: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="1B5E20")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize_columns(ws, max_width: int = 48) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        width = min(max(len(str(c.value or "")) for c in col_cells) + 2, max_width)
        ws.column_dimensions[letter].width = width


def generate_excel(users: list[tuple]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    headers = [
        "Agency Code",
        "Agency Name",
        "Email",
        "Password",
        "Role(s)",
        "Display Name",
        "Portal URL",
        "Notes",
    ]

    # Summary sheet
    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "AgriCheck V3 - Dev Users Per Agency"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    summary["A4"] = "Agency"
    summary["B4"] = "User Count"
    summary["A4"].font = summary["B4"].font = Font(bold=True)

    by_agency: dict[str, list[tuple]] = defaultdict(list)
    for user in users:
        by_agency[user[0]].append(user)

    row = 5
    for code in sorted(by_agency.keys()):
        summary.cell(row=row, column=1, value=code)
        summary.cell(row=row, column=2, value=len(by_agency[code]))
        row += 1

    summary["A" + str(row + 1)] = "Note: Agency Admin (ROLE_AGENCY_ADMIN) can manage per-agency payment config like System Admin."
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 14

    # Per-agency sheets
    for code in sorted(by_agency.keys()):
        sheet_name = code[:31]
        ws = wb.create_sheet(sheet_name)
        style_header(ws, headers)
        for r, user in enumerate(by_agency[code], 2):
            for c, value in enumerate(user, 1):
                ws.cell(row=r, column=c, value=value)
        autosize_columns(ws)

    # All Users sheet
    all_ws = wb.create_sheet("All Users")
    style_header(all_ws, headers)
    for r, user in enumerate(users, 2):
        for c, value in enumerate(user, 1):
            all_ws.cell(row=r, column=c, value=value)
    autosize_columns(all_ws)

    # Payment config sheet
    pay_ws = wb.create_sheet("Payment Config")
    pay_headers = ["Agency", "Entry Type", "Amount", "Currency", "Active", "Notes"]
    style_header(pay_ws, pay_headers)
    for r, row_data in enumerate(PAYMENT_CONFIG, 2):
        for c, value in enumerate(row_data, 1):
            pay_ws.cell(row=r, column=c, value=value)
    autosize_columns(pay_ws)

    # Role legend
    roles_ws = wb.create_sheet("Role Legend")
    role_headers = ["Role Code", "Description", "Portal"]
    style_header(roles_ws, role_headers)
    roles = [
        ("ROLE_ADMIN", "System Administrator — global payment config, agencies, users", "/admin"),
        ("ROLE_AGENCY_ADMIN", "Agency Administrator — agency-level config incl. payment fees", "/agency"),
        ("ROLE_EVALUATOR", "Document & entry evaluator", "/agency"),
        ("ROLE_BILLING_AGENT", "DA billing payment verification", "/agency"),
        ("ROLE_INSPECTOR", "BOC container inspection photos", "/inspector"),
        ("ROLE_IMPORTER", "Client / importer", "/client"),
        ("ROLE_MAV_ADMIN", "MAV program admin", "/mav/admin"),
        ("ROLE_DA_SECRETARY", "DA Secretary dashboard", "/da"),
        ("ROLE_WAREHOUSE_STAFF", "Warehouse receive", "/warehouse"),
        ("ROLE_DRIVER", "AgriTrack driver", "/driver"),
        ("ROLE_OPERATOR", "AgriTrack operator", "/driver"),
    ]
    for r, role_row in enumerate(roles, 2):
        for c, value in enumerate(role_row, 1):
            roles_ws.cell(row=r, column=c, value=value)
    autosize_columns(roles_ws)

    out = OUTPUT_DIR / "AgriCheck-V3-Users-Per-Agency.xlsx"
    wb.save(out)
    return out


class WorkflowPDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(27, 94, 32)
        self.cell(0, 8, "AgriCheck V3 - Import Entry Workflow (Approved Path)", align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} | Page {self.page_no()}", align="C")


def generate_pdf() -> Path:
    pdf = WorkflowPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(40, 40, 40)
    pdf.multi_cell(
        0,
        5,
        "Complete sequence after evaluator approves an import entry. "
        "Certificate is auto-issued when DA billing payment is verified (V3 behavior).",
    )
    pdf.ln(4)

    box_w = 190
    box_h = 22
    colors = [
        (232, 245, 233),  # green tint
        (227, 242, 253),  # blue tint
        (255, 243, 224),  # orange tint
        (237, 231, 246),  # purple tint
    ]

    for idx, (step, actor, action, status, location) in enumerate(WORKFLOW_STEPS):
        fill = colors[idx % len(colors)]
        pdf.set_fill_color(*fill)
        pdf.set_draw_color(100, 100, 100)
        y = pdf.get_y()
        if y + box_h > 270:
            pdf.add_page()
            y = pdf.get_y()

        pdf.rect(10, y, box_w, box_h, style="DF")
        pdf.set_xy(12, y + 2)
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(27, 94, 32)
        pdf.cell(0, 5, f"Step {step}: {actor}", new_x="LMARGIN", new_y="NEXT")

        pdf.set_x(12)
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(30, 30, 30)
        pdf.cell(0, 4, f"Action: {action}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_x(12)
        pdf.cell(0, 4, f"Entry Status: {status}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_x(12)
        pdf.set_font("Helvetica", "I", 8)
        pdf.set_text_color(80, 80, 80)
        pdf.cell(0, 4, f"Where: {location}", new_x="LMARGIN", new_y="NEXT")

        pdf.set_y(y + box_h + 1)
        if idx < len(WORKFLOW_STEPS) - 1:
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(27, 94, 32)
            pdf.cell(0, 5, "v", align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(180, 0, 0)
    pdf.cell(0, 6, "Alternative Outcomes (Evaluator Step 3):", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(40, 40, 40)
    pdf.multi_cell(
        0,
        5,
        "- Revision Required: ForCompliance, client revises and resubmits, back to UnderReview\n"
        "- Rejected: workflow ends",
    )

    out = OUTPUT_DIR / "AgriCheck-V3-Workflow-Approved-Path.pdf"
    pdf.output(str(out))
    return out


def main() -> None:
    users = merge_users()
    excel_path = generate_excel(users)
    pdf_path = generate_pdf()
    manifest = {
        "generated_at": datetime.now().isoformat(),
        "excel": str(excel_path),
        "pdf": str(pdf_path),
        "user_rows": len(users),
    }
    manifest_path = OUTPUT_DIR / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"Excel: {excel_path}")
    print(f"PDF:   {pdf_path}")


if __name__ == "__main__":
    main()
